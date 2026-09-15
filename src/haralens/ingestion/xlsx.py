"""Defensive, bounded XLSX ingestion from caller-owned bytes."""

import io
import logging
import math
import posixpath
import re
import time
import zipfile
import zlib
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from typing import Protocol, cast
from xml.etree.ElementTree import ParseError

import pandas as pd
from defusedxml import ElementTree as DefusedElementTree
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from haralens.ingestion._shared import (
    enforce_materialization_invariant,
    enforce_source_size,
    safe_display_name,
    validate_text_columns,
)
from haralens.ingestion.errors import (
    ArchiveSafetyError,
    ColumnLimitExceededError,
    DuplicateColumnsError,
    EmptySourceError,
    EmptyWorksheetError,
    IngestionError,
    InvalidColumnNamesError,
    MalformedWorkbookError,
    NoDataRowsError,
    NoUsableWorksheetError,
    RowLimitExceededError,
    UnsupportedWorkbookFormatError,
    UnsupportedWorksheetStructureError,
    WorkbookLimitExceededError,
    WorksheetNotFoundError,
    WorksheetSelectionRequiredError,
)
from haralens.ingestion.models import (
    ExcelFormatMetadata,
    ExcelHeaderScalar,
    IngestionMetadata,
    IngestionRequest,
    IngestionResult,
    SourceType,
)

_OLE_COMPOUND_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_XLSX_CONTENT_TYPE = b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
_MACRO_CONTENT_TYPE = b"application/vnd.ms-excel.sheet.macroenabled.main+xml"
_REQUIRED_MEMBERS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "xl/workbook.xml",
    "xl/_rels/workbook.xml.rels",
}
_DRIVE_PATH = re.compile(r"^[A-Za-z]:")


class _Cell(Protocol):
    value: object
    data_type: str


class _Worksheet(Protocol):
    title: str
    sheet_state: str
    _worksheet_path: str

    def reset_dimensions(self) -> None: ...

    def iter_rows(self, *, values_only: bool = False) -> Iterable[tuple[_Cell, ...]]: ...


class _Workbook(Protocol):
    worksheets: Sequence[_Worksheet]

    def close(self) -> None: ...


def _unsafe_archive_name(name: str) -> bool:
    portable = name.replace("\\", "/")
    normalized = posixpath.normpath(portable)
    parts = portable.split("/")
    return (
        not name
        or "\x00" in name
        or portable.startswith("/")
        or bool(_DRIVE_PATH.match(portable))
        or ".." in parts
        or normalized == ".."
        or normalized.startswith("../")
    )


def _normalized_archive_name(name: str) -> str:
    return posixpath.normpath(name.replace("\\", "/")).casefold()


def _validate_container(request: IngestionRequest) -> None:
    content = request.content
    enforce_source_size(content, maximum=request.limits.max_source_bytes)
    if not content:
        raise EmptySourceError("The XLSX source is empty (zero bytes).")
    if content.startswith(_OLE_COMPOUND_MAGIC):
        raise UnsupportedWorkbookFormatError(
            "Legacy binary Excel workbooks are not supported; provide an XLSX workbook."
        )

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > request.limits.max_xlsx_archive_entries:
                raise WorkbookLimitExceededError(
                    resource="archive entry count",
                    observed=len(entries),
                    maximum=request.limits.max_xlsx_archive_entries,
                )

            names = [entry.filename for entry in entries]
            normalized_names = [_normalized_archive_name(name) for name in names]
            if len(set(normalized_names)) != len(normalized_names):
                raise ArchiveSafetyError(
                    "The XLSX archive contains duplicate normalized member paths."
                )
            if any(_unsafe_archive_name(name) for name in names):
                raise ArchiveSafetyError("The XLSX archive contains an unsafe member path.")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ArchiveSafetyError("Encrypted XLSX archive members are not supported.")

            total_uncompressed = sum(entry.file_size for entry in entries)
            if total_uncompressed > request.limits.max_xlsx_uncompressed_bytes:
                raise WorkbookLimitExceededError(
                    resource="uncompressed archive size",
                    observed=total_uncompressed,
                    maximum=request.limits.max_xlsx_uncompressed_bytes,
                )
            for entry in entries:
                if entry.file_size == 0:
                    continue
                if entry.compress_size == 0:
                    raise ArchiveSafetyError(
                        "The XLSX archive contains an implausible compression ratio."
                    )
                ratio = entry.file_size / entry.compress_size
                if ratio > request.limits.max_xlsx_compression_ratio:
                    raise ArchiveSafetyError(
                        "The XLSX archive exceeds the configured compression-ratio limit."
                    )

            if not _REQUIRED_MEMBERS.issubset(names):
                raise UnsupportedWorkbookFormatError(
                    "The source is not a supported XLSX workbook container."
                )
            content_types = archive.read("[Content_Types].xml").lower()
            has_vba = any(name.lower().endswith("vbaproject.bin") for name in names)
            if has_vba or _MACRO_CONTENT_TYPE in content_types:
                raise UnsupportedWorkbookFormatError(
                    "Macro-enabled workbooks are not supported; provide a macro-free XLSX workbook."
                )
            if _XLSX_CONTENT_TYPE not in content_types:
                raise UnsupportedWorkbookFormatError(
                    "The source is not a supported XLSX workbook container."
                )
            if archive.testzip() is not None:
                raise MalformedWorkbookError("The XLSX archive failed its integrity check.")
    except IngestionError:
        raise
    except (zipfile.BadZipFile, zipfile.LargeZipFile, RuntimeError, zlib.error) as error:
        raise MalformedWorkbookError("The XLSX source is not a valid ZIP workbook.") from error


def _safe_worksheet_names(worksheets: Sequence[_Worksheet]) -> tuple[str, ...]:
    return tuple(
        safe_display_name(worksheet.title, fallback="worksheet") for worksheet in worksheets
    )


def _visible_usable_worksheets(
    worksheets: Sequence[_Worksheet], *, maximum_cells: int
) -> list[_Worksheet]:
    usable: list[_Worksheet] = []
    inspected_cells = 0
    for worksheet in worksheets:
        if worksheet.sheet_state != "visible":
            continue
        worksheet.reset_dimensions()
        for row in worksheet.iter_rows(values_only=False):
            inspected_cells += max(1, len(row))
            if inspected_cells > maximum_cells:
                raise WorkbookLimitExceededError(
                    resource="worksheet selection inspection cell count",
                    observed=inspected_cells,
                    maximum=maximum_cells,
                )
            if any(cell.value is not None for cell in row):
                usable.append(worksheet)
                break
    return usable


def _select_worksheet(
    workbook: _Workbook, requested_name: str | None, *, maximum_cells: int
) -> _Worksheet:
    worksheets = workbook.worksheets
    if not worksheets:
        raise NoUsableWorksheetError("The workbook contains no worksheets.")

    if requested_name is not None:
        for worksheet in worksheets:
            if worksheet.title == requested_name:
                return worksheet
        raise WorksheetNotFoundError(available_worksheets=_safe_worksheet_names(worksheets))

    usable = _visible_usable_worksheets(worksheets, maximum_cells=maximum_cells)
    if not usable:
        raise NoUsableWorksheetError(
            "The workbook contains no visible worksheet with any values; "
            "select a populated hidden worksheet explicitly if needed."
        )
    if len(usable) > 1:
        raise WorksheetSelectionRequiredError(available_worksheets=_safe_worksheet_names(usable))
    return usable[0]


def _header_has_merge(content: bytes, worksheet_path: str, header_row: int) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            with archive.open(worksheet_path) as worksheet_xml:
                for _, element in DefusedElementTree.iterparse(worksheet_xml, events=("end",)):
                    if element.tag.rsplit("}", 1)[-1] == "mergeCell":
                        reference = element.attrib.get("ref", "")
                        try:
                            _, minimum_row, _, maximum_row = range_boundaries(reference)
                        except ValueError as error:
                            raise MalformedWorkbookError(
                                "The XLSX worksheet contains an invalid merged-cell range."
                            ) from error
                        if minimum_row is None or maximum_row is None:
                            raise MalformedWorkbookError(
                                "The XLSX worksheet contains an invalid merged-cell range."
                            )
                        if minimum_row <= header_row <= maximum_row:
                            return True
                    element.clear()
    except IngestionError:
        raise
    except (KeyError, DefusedXmlException, ParseError, zipfile.BadZipFile, OSError) as error:
        raise MalformedWorkbookError("The XLSX worksheet structure is malformed.") from error
    return False


def _trim_trailing_empty(row: tuple[object, ...]) -> tuple[object, ...]:
    final = len(row)
    while final and row[final - 1] is None:
        final -= 1
    return row[:final]


def _canonicalize_header(
    values: tuple[object, ...], data_types: tuple[str, ...]
) -> tuple[list[str], tuple[ExcelHeaderScalar, ...]]:
    canonical: list[str] = []
    originals: list[ExcelHeaderScalar] = []
    invalid_positions: list[int] = []

    for position, (value, data_type) in enumerate(zip(values, data_types, strict=True), start=1):
        name: str | None = None
        if data_type == "e" or value is None:
            pass
        elif isinstance(value, str):
            if value.strip():
                name = value
        elif isinstance(value, bool):
            name = "TRUE" if value else "FALSE"
        elif isinstance(value, int):
            name = str(value)
        elif isinstance(value, float):
            if math.isfinite(value):
                name = repr(value)
        elif isinstance(value, datetime):
            name = value.isoformat()
        elif isinstance(value, date):
            name = value.isoformat()

        if name is None:
            invalid_positions.append(position)
            continue
        canonical.append(name)
        originals.append(cast(ExcelHeaderScalar, value))

    if invalid_positions:
        raise InvalidColumnNamesError(column_positions=tuple(invalid_positions))

    seen_originals: set[tuple[type[object], ExcelHeaderScalar]] = set()
    structural_duplicates: list[str] = []
    for original, name in zip(originals, canonical, strict=True):
        key = (type(original), original)
        if key in seen_originals and name not in structural_duplicates:
            structural_duplicates.append(name)
        seen_originals.add(key)
    if structural_duplicates:
        raise DuplicateColumnsError(duplicate_columns=tuple(structural_duplicates))

    return validate_text_columns(canonical), tuple(originals)


def _read_worksheet(
    request: IngestionRequest, worksheet: _Worksheet
) -> tuple[pd.DataFrame, tuple[ExcelHeaderScalar, ...]]:
    worksheet.reset_dimensions()
    header: tuple[object, ...] | None = None
    header_data_types: tuple[str, ...] | None = None
    header_row = 0
    data_rows: list[tuple[object, ...]] = []
    pending_blank_rows = 0
    scanned_cells = 0
    width = 0

    for physical_row, cells in enumerate(worksheet.iter_rows(values_only=False), start=1):
        scanned_cells += max(1, len(cells))
        if scanned_cells > request.limits.max_xlsx_cells:
            raise WorkbookLimitExceededError(
                resource="scanned worksheet cell count",
                observed=scanned_cells,
                maximum=request.limits.max_xlsx_cells,
            )

        raw_row = tuple(cell.value for cell in cells)
        row = _trim_trailing_empty(raw_row)
        if len(row) > request.limits.max_columns:
            raise ColumnLimitExceededError(
                observed_columns=len(row), max_columns=request.limits.max_columns
            )
        if header is None:
            if not row:
                continue
            header = row
            header_data_types = tuple(cell.data_type for cell in cells[: len(row)])
            header_row = physical_row
            width = len(row)
            continue

        if not row:
            pending_blank_rows += 1
            continue

        if len(data_rows) + pending_blank_rows + 1 > request.limits.max_rows:
            raise RowLimitExceededError(
                observed_rows=len(data_rows) + pending_blank_rows + 1,
                max_rows=request.limits.max_rows,
            )
        data_rows.extend([tuple()] * pending_blank_rows)
        pending_blank_rows = 0
        data_rows.append(row)
        width = max(width, len(row))

    if header is None or header_data_types is None:
        raise EmptyWorksheetError("The selected worksheet contains no values.")
    if not data_rows:
        raise NoDataRowsError("The selected worksheet contains a header but no data rows.")
    if (len(data_rows) + 1) * width > request.limits.max_xlsx_cells:
        raise WorkbookLimitExceededError(
            resource="materialized worksheet cell count",
            observed=(len(data_rows) + 1) * width,
            maximum=request.limits.max_xlsx_cells,
        )

    if _header_has_merge(request.content, worksheet._worksheet_path, header_row):
        raise UnsupportedWorksheetStructureError(
            "Merged cells intersect the worksheet header and are not supported."
        )
    padded_header = (*header, *([None] * (width - len(header))))
    padded_data_types = (*header_data_types, *(["n"] * (width - len(header_data_types))))
    columns, original_headers = _canonicalize_header(padded_header, padded_data_types)

    normalized_rows = [list(row) + [None] * (width - len(row)) for row in data_rows]
    table = pd.DataFrame(normalized_rows, columns=columns)
    enforce_materialization_invariant(table, expected_rows=len(data_rows), expected_columns=columns)
    return table, original_headers


class XlsxIngestionAdapter:
    """Load macro-free XLSX workbooks with ZIP and worksheet bounds."""

    source_type = SourceType.XLSX

    def __init__(self) -> None:
        self._logger = logging.getLogger("haralens.ingestion.xlsx")

    def ingest(self, request: IngestionRequest) -> IngestionResult:
        started = time.perf_counter()
        safe_name = safe_display_name(request.source_name, fallback="unnamed.xlsx")
        source_bytes = len(request.content)
        log_data: dict[str, object] = {
            "source_type": self.source_type.value,
            "source_name": safe_name,
            "source_bytes": source_bytes,
        }
        self._logger.info("ingestion_started", extra={"event_data": log_data})
        try:
            table, format_metadata = self._load(request)
        except IngestionError as error:
            duration_ms = (time.perf_counter() - started) * 1_000
            self._logger.warning(
                "ingestion_failed",
                extra={
                    "event_data": {
                        **log_data,
                        "error_code": error.code.value,
                        "duration_ms": round(duration_ms, 3),
                    }
                },
            )
            raise

        duration_ms = (time.perf_counter() - started) * 1_000
        metadata = IngestionMetadata(
            source_type=self.source_type,
            source_name=safe_name,
            source_bytes=source_bytes,
            encoding=None,
            duration_ms=duration_ms,
            format_metadata=format_metadata,
        )
        self._logger.info(
            "ingestion_succeeded",
            extra={
                "event_data": {
                    **log_data,
                    "selected_worksheet": format_metadata.selected_worksheet,
                    "row_count": len(table.index),
                    "column_count": len(table.columns),
                    "duration_ms": round(duration_ms, 3),
                }
            },
        )
        return IngestionResult(
            table=table,
            row_count=len(table.index),
            column_count=len(table.columns),
            metadata=metadata,
        )

    def _load(self, request: IngestionRequest) -> tuple[pd.DataFrame, ExcelFormatMetadata]:
        _validate_container(request)
        workbook: _Workbook | None = None
        try:
            workbook = cast(
                _Workbook,
                load_workbook(
                    io.BytesIO(request.content),
                    read_only=True,
                    data_only=False,
                    keep_vba=False,
                    keep_links=False,
                ),
            )
            worksheet_count = len(workbook.worksheets)
            if worksheet_count > request.limits.max_xlsx_worksheets:
                raise WorkbookLimitExceededError(
                    resource="worksheet count",
                    observed=worksheet_count,
                    maximum=request.limits.max_xlsx_worksheets,
                )
            worksheet = _select_worksheet(
                workbook,
                request.worksheet_name,
                maximum_cells=request.limits.max_xlsx_cells,
            )
            table, original_headers = _read_worksheet(request, worksheet)
            visible_count = sum(sheet.sheet_state == "visible" for sheet in workbook.worksheets)
            return table, ExcelFormatMetadata(
                selected_worksheet=worksheet.title,
                worksheet_count=worksheet_count,
                visible_worksheet_count=visible_count,
                original_column_headers=original_headers,
            )
        except IngestionError:
            raise
        except (
            InvalidFileException,
            DefusedXmlException,
            ParseError,
            KeyError,
            OSError,
            TypeError,
            ValueError,
        ) as error:
            raise MalformedWorkbookError("The XLSX workbook could not be parsed safely.") from error
        finally:
            if workbook is not None:
                workbook.close()
