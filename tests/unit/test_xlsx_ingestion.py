import io
import json
import struct
import zipfile
from datetime import date, datetime, time

import pandas as pd
import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart

from haralens.common.logging import configure_logging
from haralens.ingestion.errors import (
    ArchiveSafetyError,
    ColumnLimitExceededError,
    DuplicateColumnsError,
    EmptyWorksheetError,
    IngestionConsistencyError,
    InvalidColumnNamesError,
    MalformedWorkbookError,
    NoDataRowsError,
    NoUsableWorksheetError,
    RowLimitExceededError,
    SourceTooLargeError,
    UnsupportedWorkbookFormatError,
    UnsupportedWorksheetStructureError,
    WorkbookLimitExceededError,
    WorksheetNotFoundError,
    WorksheetSelectionRequiredError,
)
from haralens.ingestion.models import (
    ExcelFormatMetadata,
    IngestionRequest,
    ResourceLimits,
    SourceType,
)
from haralens.ingestion.xlsx import (
    XlsxIngestionAdapter,
    _canonicalize_header,
    _header_has_merge,
    _trim_trailing_empty,
)

DEFAULT_LIMITS = ResourceLimits(max_source_bytes=1_000_000, max_rows=100, max_columns=20)


def workbook_bytes(
    sheets: dict[str, list[list[object]]],
    *,
    hidden: set[str] | None = None,
    very_hidden: set[str] | None = None,
    merges: dict[str, list[str]] | None = None,
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        worksheet = workbook.create_sheet(name)
        for row in rows:
            worksheet.append(row)
        if hidden and name in hidden:
            worksheet.sheet_state = "hidden"
        if very_hidden and name in very_hidden:
            worksheet.sheet_state = "veryHidden"
        for cell_range in (merges or {}).get(name, []):
            worksheet.merge_cells(cell_range)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def rewrite_zip(
    content: bytes,
    *,
    replacements: dict[str, bytes] | None = None,
    additions: dict[str, bytes] | None = None,
    removals: set[str] | None = None,
) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source:
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as destination:
            for entry in source.infolist():
                if removals and entry.filename in removals:
                    continue
                value = source.read(entry.filename)
                destination.writestr(
                    entry.filename, (replacements or {}).get(entry.filename, value)
                )
            for name, value in (additions or {}).items():
                destination.writestr(name, value)
    return output.getvalue()


def mark_first_central_entry_encrypted(content: bytes) -> bytes:
    changed = bytearray(content)
    central_offset = changed.index(b"PK\x01\x02")
    flags = struct.unpack_from("<H", changed, central_offset + 8)[0]
    struct.pack_into("<H", changed, central_offset + 8, flags | 0x1)
    return bytes(changed)


def mark_first_central_entry_compressed_size_zero(content: bytes) -> bytes:
    changed = bytearray(content)
    central_offset = changed.index(b"PK\x01\x02")
    struct.pack_into("<I", changed, central_offset + 20, 0)
    return bytes(changed)


def corrupt_member_payload(content: bytes, member_name: str) -> bytes:
    changed = bytearray(content)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entry = archive.getinfo(member_name)
    filename_length, extra_length = struct.unpack_from("<HH", changed, entry.header_offset + 26)
    payload_offset = entry.header_offset + 30 + filename_length + extra_length
    changed[payload_offset + entry.compress_size // 2] ^= 0xFF
    return bytes(changed)


def ingest(
    content: bytes,
    *,
    source_name: str = "data.xlsx",
    limits: ResourceLimits = DEFAULT_LIMITS,
    worksheet_name: str | None = None,
):
    return XlsxIngestionAdapter().ingest(
        IngestionRequest(
            content=content,
            source_name=source_name,
            limits=limits,
            worksheet_name=worksheet_name,
        )
    )


def test_valid_one_sheet_xlsx_returns_values_and_metadata() -> None:
    value_date = datetime(2025, 1, 2, 3, 4, 5)
    content = workbook_bytes(
        {"Data": [["name", "amount", "active", "when"], ["A", 3.5, True, value_date]]}
    )

    result = ingest(content)

    assert result.row_count == 1
    assert result.column_count == 4
    assert result.table.iloc[0].tolist() == ["A", 3.5, True, value_date]
    assert result.metadata.source_type is SourceType.XLSX
    assert result.metadata.encoding is None
    assert result.metadata.source_bytes == len(content)
    assert isinstance(result.metadata.format_metadata, ExcelFormatMetadata)
    assert result.metadata.format_metadata.selected_worksheet == "Data"
    assert result.metadata.format_metadata.formula_representation == "formula_text"


def test_multiple_visible_sheets_require_an_explicit_selection() -> None:
    content = workbook_bytes({"First": [["a"], [1]], "Second": [["b"], [2]]})
    with pytest.raises(WorksheetSelectionRequiredError) as raised:
        ingest(content)
    assert raised.value.available_worksheets == ("First", "Second")


@pytest.mark.parametrize("empty_sheet_count", [1, 3])
def test_empty_visible_sheets_do_not_force_selection(empty_sheet_count: int) -> None:
    sheets: dict[str, list[list[object]]] = {"Data": [["a"], [1]]}
    sheets.update({f"Empty {index}": [] for index in range(empty_sheet_count)})

    result = ingest(workbook_bytes(sheets))

    assert result.metadata.format_metadata.selected_worksheet == "Data"


def test_all_visible_sheets_empty_have_no_usable_worksheet() -> None:
    content = workbook_bytes({"Empty One": [], "Empty Two": []})
    with pytest.raises(NoUsableWorksheetError, match="no visible worksheet with any values"):
        ingest(content)


def test_empty_visible_sheet_does_not_expose_populated_hidden_sheet_automatically() -> None:
    content = workbook_bytes({"Empty": [], "Hidden Data": [["a"], [1]]}, hidden={"Hidden Data"})
    with pytest.raises(NoUsableWorksheetError):
        ingest(content)
    assert ingest(content, worksheet_name="Hidden Data").row_count == 1


def test_explicit_empty_sheet_selection_retains_empty_worksheet_failure() -> None:
    content = workbook_bytes({"Data": [["a"], [1]], "Empty": []})
    with pytest.raises(EmptyWorksheetError):
        ingest(content, worksheet_name="Empty")


def test_explicit_sheet_selection_is_exact_and_deterministic() -> None:
    content = workbook_bytes({"First": [["a"], [1]], "Second": [["b"], [2]]})
    result = ingest(content, worksheet_name="Second")
    assert list(result.table.columns) == ["b"]
    assert result.table.iloc[0, 0] == 2

    with pytest.raises(WorksheetNotFoundError) as raised:
        ingest(content, worksheet_name="second")
    assert raised.value.available_worksheets == ("First", "Second")


def test_hidden_sheets_are_not_automatic_but_can_be_selected() -> None:
    content = workbook_bytes(
        {"Visible": [["a"], [1]], "Hidden": [["secret"], [2]]}, hidden={"Hidden"}
    )
    assert ingest(content).metadata.format_metadata.selected_worksheet == "Visible"
    assert ingest(content, worksheet_name="Hidden").table.iloc[0, 0] == 2


def test_very_hidden_sheet_requires_explicit_selection() -> None:
    content = workbook_bytes(
        {"Visible": [["a"], [1]], "Internal": [["b"], [2]]},
        very_hidden={"Internal"},
    )
    assert ingest(content).metadata.format_metadata.selected_worksheet == "Visible"
    assert ingest(content, worksheet_name="Internal").table.iloc[0, 0] == 2


def test_workbook_with_no_visible_sheet_fails_clearly() -> None:
    content = workbook_bytes({"Only": [["a"], [1]]})
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        workbook_xml = archive.read("xl/workbook.xml")
    hidden_xml = workbook_xml.replace(b'state="visible"', b'state="hidden"')
    content = rewrite_zip(content, replacements={"xl/workbook.xml": hidden_xml})

    with pytest.raises(NoUsableWorksheetError):
        ingest(content)
    assert ingest(content, worksheet_name="Only").row_count == 1


def test_chart_only_workbook_has_no_usable_worksheet() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    chart_sheet = workbook.create_chartsheet("Chart")
    chart_sheet.add_chart(BarChart())
    workbook.remove(worksheet)
    output = io.BytesIO()
    workbook.save(output)

    with pytest.raises(NoUsableWorksheetError, match="no worksheets"):
        ingest(output.getvalue())


@pytest.mark.parametrize(
    ("rows", "error_type"),
    [([], EmptyWorksheetError), ([["a"]], NoDataRowsError)],
)
def test_empty_and_header_only_worksheets_are_distinct(
    rows: list[list[object]], error_type: type[Exception]
) -> None:
    with pytest.raises(error_type):
        ingest(workbook_bytes({"Data": rows}), worksheet_name="Data")


def test_duplicate_empty_and_whitespace_headers_are_rejected() -> None:
    with pytest.raises(DuplicateColumnsError):
        ingest(workbook_bytes({"Data": [["a", "a"], [1, 2]]}))
    with pytest.raises(InvalidColumnNamesError) as empty:
        ingest(workbook_bytes({"Data": [[None, "b"], [1, 2]]}))
    assert empty.value.column_positions == (1,)
    with pytest.raises(InvalidColumnNamesError) as whitespace:
        ingest(workbook_bytes({"Data": [["a", "  "], [1, 2]]}))
    assert whitespace.value.column_positions == (2,)


def test_unusual_valid_headers_are_preserved_exactly() -> None:
    headers = [" customer id ", "Total (€)", "a.b", "日本語"]
    result = ingest(workbook_bytes({"Data": [headers, [1, 2, 3, "四"]]}))
    assert list(result.table.columns) == headers


def test_safe_scalar_headers_are_canonicalized_deterministically() -> None:
    header_date = date(2024, 1, 2)
    header_datetime = datetime(2024, 1, 2, 3, 4, 5)
    columns, originals = _canonicalize_header(
        (" exact ", 2024, 1.5, header_date, header_datetime, True, False, "=1+1"),
        ("s", "n", "n", "d", "d", "b", "b", "f"),
    )
    assert columns == [
        " exact ",
        "2024",
        "1.5",
        "2024-01-02",
        "2024-01-02T03:04:05",
        "TRUE",
        "FALSE",
        "=1+1",
    ]
    assert originals == (
        " exact ",
        2024,
        1.5,
        header_date,
        header_datetime,
        True,
        False,
        "=1+1",
    )


def test_numeric_boolean_and_formula_headers_work_end_to_end() -> None:
    result = ingest(workbook_bytes({"Data": [[2024, 1.5, True, "=1+1"], [1, 2, 3, 4]]}))
    assert list(result.table.columns) == ["2024", "1.5", "TRUE", "=1+1"]
    assert result.metadata.format_metadata.original_column_headers == (2024, 1.5, True, "=1+1")
    assert "original_column_headers" in result.metadata.model_dump_json()


def test_excel_date_and_datetime_headers_use_loaded_iso_representation() -> None:
    result = ingest(
        workbook_bytes({"Data": [[date(2024, 1, 2), datetime(2024, 1, 2, 3, 4, 5)], [1, 2]]})
    )
    assert list(result.table.columns) == ["2024-01-02T00:00:00", "2024-01-02T03:04:05"]
    assert result.metadata.format_metadata.original_column_headers == (
        datetime(2024, 1, 2),
        datetime(2024, 1, 2, 3, 4, 5),
    )


@pytest.mark.parametrize(
    ("values", "data_types", "positions"),
    [
        ((None,), ("n",), (1,)),
        (("   ",), ("s",), (1,)),
        ((float("nan"), float("inf")), ("n", "n"), (1, 2)),
        ((time(12, 30),), ("d",), (1,)),
        (("#DIV/0!",), ("e",), (1,)),
    ],
)
def test_unsafe_or_unsupported_scalar_headers_are_rejected(
    values: tuple[object, ...], data_types: tuple[str, ...], positions: tuple[int, ...]
) -> None:
    with pytest.raises(InvalidColumnNamesError) as raised:
        _canonicalize_header(values, data_types)
    assert raised.value.column_positions == positions


def test_excel_error_header_is_rejected_end_to_end() -> None:
    with pytest.raises(InvalidColumnNamesError) as raised:
        ingest(workbook_bytes({"Data": [["#DIV/0!"], [1]]}))
    assert raised.value.column_positions == (1,)


@pytest.mark.parametrize(
    "headers",
    [
        [2024, 2024],
        [2024, "2024"],
        [1.5, "1.5"],
        [True, "TRUE"],
    ],
)
def test_structural_duplicates_and_canonicalization_collisions_are_rejected(
    headers: list[object],
) -> None:
    with pytest.raises(DuplicateColumnsError):
        ingest(workbook_bytes({"Data": [headers, [1, 2]]}))


def test_row_and_column_limits_are_enforced_before_dataframe_creation() -> None:
    content = workbook_bytes({"Data": [["a"], [1], [2]]})
    with pytest.raises(RowLimitExceededError):
        ingest(
            content,
            limits=ResourceLimits(max_source_bytes=len(content), max_rows=1, max_columns=1),
        )

    wide = workbook_bytes({"Data": [["a", "b"], [1, 2]]})
    with pytest.raises(ColumnLimitExceededError):
        ingest(
            wide,
            limits=ResourceLimits(max_source_bytes=len(wide), max_rows=1, max_columns=1),
        )


def test_source_byte_limit_accepts_boundary_and_rejects_above_it() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    limits = ResourceLimits(max_source_bytes=len(content), max_rows=1, max_columns=1)
    assert ingest(content, limits=limits).row_count == 1
    with pytest.raises(SourceTooLargeError):
        ingest(
            content,
            limits=ResourceLimits(max_source_bytes=len(content) - 1, max_rows=1, max_columns=1),
        )


@pytest.mark.parametrize("content", [b"not an xlsx", b"PK\x03\x04broken"])
def test_malformed_and_fake_xlsx_sources_fail_safely(content: bytes) -> None:
    with pytest.raises(MalformedWorkbookError):
        ingest(content)


def test_zero_byte_xlsx_is_a_typed_empty_source() -> None:
    from haralens.ingestion.errors import EmptySourceError

    with pytest.raises(EmptySourceError):
        ingest(b"")


def test_legacy_and_macro_enabled_workbooks_are_typed_unsupported_formats() -> None:
    with pytest.raises(UnsupportedWorkbookFormatError):
        ingest(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"legacy")

    content = workbook_bytes({"Data": [["a"], [1]]})
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        content_types = archive.read("[Content_Types].xml")
    macro_types = content_types.replace(
        b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    )
    macro_content = rewrite_zip(content, replacements={"[Content_Types].xml": macro_types})
    with pytest.raises(UnsupportedWorkbookFormatError):
        ingest(macro_content)

    vba_content = rewrite_zip(content, additions={"xl/vbaProject.bin": b"macro"})
    with pytest.raises(UnsupportedWorkbookFormatError):
        ingest(vba_content)


def test_malformed_workbook_xml_is_translated_without_parser_details() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        workbook_xml = archive.read("xl/workbook.xml")
    malformed = rewrite_zip(
        content,
        replacements={"xl/workbook.xml": workbook_xml.replace(b"</workbook>", b"")},
    )
    with pytest.raises(MalformedWorkbookError) as raised:
        ingest(malformed)
    assert "XML" not in raised.value.message
    assert raised.value.__cause__ is not None


@pytest.mark.parametrize(
    "unsafe_name",
    ["../secret.txt", "..\\secret.txt", "xl/../../secret.txt", "/absolute.txt", "C:\\x.txt"],
)
def test_archive_path_traversal_is_portable(unsafe_name: str) -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    unsafe = rewrite_zip(content, additions={unsafe_name: b"secret"})
    with pytest.raises(ArchiveSafetyError, match="unsafe member path"):
        ingest(unsafe)


def test_duplicate_normalized_archive_paths_are_rejected() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    duplicate = rewrite_zip(content, additions={"xl/./workbook.xml": b"duplicate"})
    with pytest.raises(ArchiveSafetyError, match="duplicate normalized"):
        ingest(duplicate)


def test_encrypted_member_flag_is_rejected_before_member_reads() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    with pytest.raises(ArchiveSafetyError, match="Encrypted"):
        ingest(mark_first_central_entry_encrypted(content))


def test_positive_member_with_zero_compressed_size_is_rejected() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    impossible_ratio = mark_first_central_entry_compressed_size_zero(content)
    with pytest.raises(ArchiveSafetyError, match="implausible compression ratio"):
        ingest(impossible_ratio)


def test_missing_members_and_wrong_content_type_are_unsupported() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    missing = rewrite_zip(content, removals={"xl/_rels/workbook.xml.rels"})
    with pytest.raises(UnsupportedWorkbookFormatError):
        ingest(missing)

    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        content_types = archive.read("[Content_Types].xml")
    wrong_type = rewrite_zip(
        content,
        replacements={
            "[Content_Types].xml": content_types.replace(
                b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                b"application/octet-stream",
            )
        },
    )
    with pytest.raises(UnsupportedWorkbookFormatError):
        ingest(wrong_type)


def test_crc_corruption_is_a_typed_malformed_workbook() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    corrupted = corrupt_member_payload(content, "docProps/core.xml")
    with pytest.raises(MalformedWorkbookError):
        ingest(corrupted)


def test_empty_archive_member_is_allowed_and_integrity_checked() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    content = rewrite_zip(content, additions={"docProps/empty.bin": b""})
    assert ingest(content).row_count == 1


def test_archive_guards_run_before_decompression_or_workbook_parsing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})

    def forbidden(*args: object, **kwargs: object) -> None:
        pytest.fail("decompression or workbook parsing ran before metadata guards")

    monkeypatch.setattr(zipfile.ZipFile, "read", forbidden)
    monkeypatch.setattr(zipfile.ZipFile, "testzip", forbidden)
    monkeypatch.setattr("haralens.ingestion.xlsx.load_workbook", forbidden)
    with pytest.raises(ArchiveSafetyError, match="compression-ratio"):
        ingest(
            content,
            limits=ResourceLimits(
                max_source_bytes=len(content),
                max_rows=1,
                max_columns=1,
                max_xlsx_compression_ratio=1.0,
            ),
        )


def test_archive_entry_uncompressed_worksheet_and_cell_limits_are_enforced() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        uncompressed = sum(entry.file_size for entry in entries)

    with pytest.raises(WorkbookLimitExceededError) as entry_limit:
        ingest(
            content,
            limits=ResourceLimits(
                max_source_bytes=len(content),
                max_rows=1,
                max_columns=1,
                max_xlsx_archive_entries=len(entries) - 1,
            ),
        )
    assert entry_limit.value.resource == "archive entry count"

    with pytest.raises(WorkbookLimitExceededError) as size_limit:
        ingest(
            content,
            limits=ResourceLimits(
                max_source_bytes=len(content),
                max_rows=1,
                max_columns=1,
                max_xlsx_uncompressed_bytes=uncompressed - 1,
            ),
        )
    assert size_limit.value.resource == "uncompressed archive size"

    many_sheets = workbook_bytes({"One": [["a"], [1]], "Two": [["b"], [2]]})
    with pytest.raises(WorkbookLimitExceededError) as sheet_limit:
        ingest(
            many_sheets,
            worksheet_name="One",
            limits=ResourceLimits(
                max_source_bytes=len(many_sheets),
                max_rows=1,
                max_columns=1,
                max_xlsx_worksheets=1,
            ),
        )
    assert sheet_limit.value.resource == "worksheet count"

    with pytest.raises(WorkbookLimitExceededError) as cell_limit:
        ingest(
            content,
            limits=ResourceLimits(
                max_source_bytes=len(content),
                max_rows=1,
                max_columns=1,
                max_xlsx_cells=1,
            ),
        )
    assert "cell count" in cell_limit.value.resource


def test_auto_selection_probe_uses_the_workbook_cell_budget() -> None:
    workbook = Workbook()
    empty = workbook.active
    empty.title = "Styled Empty"
    for row in range(1, 8):
        empty.cell(row=row, column=1).number_format = "0.00"
    populated = workbook.create_sheet("Data")
    populated.append(["a"])
    populated.append([1])
    output = io.BytesIO()
    workbook.save(output)
    content = output.getvalue()

    with pytest.raises(WorkbookLimitExceededError) as raised:
        ingest(
            content,
            limits=ResourceLimits(
                max_source_bytes=len(content),
                max_rows=1,
                max_columns=1,
                max_xlsx_cells=5,
            ),
        )
    assert raised.value.resource == "worksheet selection inspection cell count"


def test_dense_materialization_cell_limit_catches_sparse_wide_rows() -> None:
    content = workbook_bytes(
        {"Data": [["a"], [1], [None, None, None, None, None, None, None, None, None, 2]]}
    )
    with pytest.raises(WorkbookLimitExceededError) as raised:
        ingest(
            content,
            limits=ResourceLimits(
                max_source_bytes=len(content),
                max_rows=2,
                max_columns=10,
                max_xlsx_cells=20,
            ),
        )
    assert raised.value.resource == "materialized worksheet cell count"


def test_formulas_are_returned_as_text_and_never_calculated() -> None:
    result = ingest(workbook_bytes({"Data": [["formula"], ["=2+2"]]}))
    assert result.table.iloc[0, 0] == "=2+2"


def test_merged_header_is_rejected_and_merged_data_is_deterministic() -> None:
    header_merge = workbook_bytes(
        {"Data": [["combined", None], [1, 2]]}, merges={"Data": ["A1:B1"]}
    )
    with pytest.raises(UnsupportedWorksheetStructureError):
        ingest(header_merge)

    data_merge = workbook_bytes({"Data": [["a", "b"], [1, 2], [3, 4]]}, merges={"Data": ["A2:B2"]})
    result = ingest(data_merge)
    assert result.table.iloc[0].tolist()[0] == 1
    assert pd.isna(result.table.iloc[0, 1])


@pytest.mark.parametrize("bad_reference", [b"A:A", b"INVALID"])
def test_malformed_merged_header_range_fails_safely(bad_reference: bytes) -> None:
    content = workbook_bytes({"Data": [["combined", None], [1, 2]]}, merges={"Data": ["A1:B1"]})
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        worksheet_xml = archive.read("xl/worksheets/sheet1.xml")
    malformed = rewrite_zip(
        content,
        replacements={
            "xl/worksheets/sheet1.xml": worksheet_xml.replace(
                b'ref="A1:B1"', b'ref="' + bad_reference + b'"'
            )
        },
    )
    with pytest.raises(MalformedWorkbookError):
        _header_has_merge(malformed, "xl/worksheets/sheet1.xml", 1)
    with pytest.raises(MalformedWorkbookError):
        ingest(malformed)


def test_missing_worksheet_member_is_translated_by_merge_inspection() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    with pytest.raises(MalformedWorkbookError):
        _header_has_merge(content, "xl/worksheets/missing.xml", 1)


def test_trailing_empty_cells_are_removed_before_width_validation() -> None:
    assert _trim_trailing_empty(("a", None, None)) == ("a",)


def test_leading_and_internal_empty_rows_have_documented_behavior() -> None:
    result = ingest(workbook_bytes({"Data": [[], ["a"], [1], [], [2], []]}))
    assert result.table["a"].tolist()[0] == 1
    assert pd.isna(result.table["a"].tolist()[1])
    assert result.table["a"].tolist()[2] == 2


def test_source_name_is_portably_sanitized() -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    result = ingest(content, source_name="../windows\\odd\n\tname.xlsx")
    assert result.metadata.source_name == "odd__name.xlsx"


def test_materialization_invariant_rejects_dataframe_shape_change(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    content = workbook_bytes({"Data": [["a"], [1]]})
    changed = pd.DataFrame({"a": [1, 2]})
    monkeypatch.setattr(
        "haralens.ingestion.xlsx.pd.DataFrame",
        lambda *args, **kwargs: changed,
    )
    with pytest.raises(IngestionConsistencyError):
        ingest(content)


def test_logs_exclude_cell_and_formula_values(capsys: pytest.CaptureFixture[str]) -> None:
    configure_logging("INFO")
    content = workbook_bytes({"Data": [["secret"], ["DO_NOT_LOG"]]})
    ingest(content, source_name="folder/private.xlsx")
    records = [json.loads(line) for line in capsys.readouterr().err.splitlines()]
    assert records[-1]["selected_worksheet"] == "Data"
    assert records[-1]["source_name"] == "private.xlsx"
    assert "DO_NOT_LOG" not in json.dumps(records)
